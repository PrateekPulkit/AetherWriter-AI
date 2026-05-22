import docx
import os

def analyze_docx(file_path):
    if not os.path.exists(file_path):
        print(f"File {file_path} not found.")
        return

    doc = docx.Document(file_path)
    
    with open('template_analysis.md', 'w', encoding='utf-8') as f:
        f.write("# Template Analysis: AI_REPORT_FINAL.docx\n\n")
        
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                style = p.style.name
                text = p.text.strip()
                
                # Identify potential headings
                if style.startswith('Heading') or text.isupper():
                    f.write(f"\n## {text} (Style: {style})\n")
                else:
                    # Just sample some text to get length/content type
                    if len(text) > 200:
                        f.write(f"\n[Long Paragraph - ~{len(text)} chars]\n")
                    else:
                        f.write(f"- {text}\n")
        
        # Check for images/diagrams placeholders
        f.write("\n\n## Potential Diagrams/Images\n")
        # In docx, images are usually in inline_shapes
        for shape in doc.inline_shapes:
            f.write(f"- Image found: {shape.type}\n")

if __name__ == "__main__":
    analyze_docx('AI_REPORT_FINAL.docx')
